#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Проверки массового расчёта: обязательная дата платежа, поведение сверки,
лист «Сверка». Ищет колонки по названию, чтобы тесты не ломались при
изменении состава отчёта.
"""

from io import BytesIO
import openpyxl
from app import app, init_db

init_db()


def make_book(charges, payments, with_payment_date=True):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Начисления"
    ws.append(["Лицевой счёт", "Период", "Начисление", "Срок оплаты", "Вид долга"])
    for c in charges:
        ws.append(list(c))
    ws2 = wb.create_sheet("Платежи")
    ws2.append(["Лицевой счёт", "Дата платежа", "Сумма платежа"])
    for p in payments:
        ws2.append(list(p))
    b = BytesIO()
    wb.save(b)
    return b.getvalue()


def sheet_rows(wb, name, header_row=1):
    """Возвращает список словарей: заголовок -> значение."""
    ws = wb[name]
    rows = list(ws.iter_rows(values_only=True))
    if header_row > len(rows):
        return []
    header = rows[header_row - 1]
    out = []
    for r in rows[header_row:]:
        if not r or all(v is None for v in r):
            continue
        out.append({h: v for h, v in zip(header, r) if h})
    return out


def find(row, *fragments):
    """Ищет значение по фрагменту названия колонки."""
    for key, value in row.items():
        low = str(key).lower()
        if all(f.lower() in low for f in fragments):
            return value
    return None


with app.test_client() as c:
    c.post('/login', data={'username': 'admin', 'password': 'admin123'})

    print("=== 1. Корректный реестр: лист «Сверка», статус «нарушений нет» ===")
    raw = make_book([("ЛС-0001", "01.01.2026", 2500, "10.02.2026", "ЖКУ"),
                     ("ЛС-0001", "01.02.2026", 2600, "15.03.2026", "ЖКУ")],
                    [("ЛС-0001", "15.04.2026", 1000), ("ЛС-0001", "05.06.2026", 1500)])
    r = c.post('/batch_calculate',
               data={'file': (BytesIO(raw), 'r.xlsx'), 'calc_date': '2026-10-04',
                     'rate_mode': 'fixed95'},
               content_type='multipart/form-data')
    assert r.status_code == 200, r.data.decode()[:300]
    wb = openpyxl.load_workbook(BytesIO(r.data), data_only=True)
    print("  листы:", wb.sheetnames)
    assert "Сверка" in wb.sheetnames, "нет листа «Сверка»"
    for row in sheet_rows(wb, "Сверка", header_row=4):
        if row.get("Лицевой счёт"):
            print(f"  Сверка: {find(row, 'лицевой')} | начислено {find(row, 'начислено')} "
                  f"| оплачено {find(row, 'оплачено')} | статус {find(row, 'статус')}")
            assert find(row, "статус") == "нарушений нет"
    # остаток после платежа виден на листе «Платежи»
    pays = [row for row in sheet_rows(wb, "Платежи") if row.get("Лицевой счёт")]
    print(f"  строк платежей: {len(pays)}; остаток после 2-го: "
          f"{find(pays[-1], 'остаток')}")

    print()
    print("=== 2. Платёж без даты — расчёт остановлен ===")
    raw2 = make_book([("ЛС-0001", "01.02.2026", 2600, "15.03.2026", "ЖКУ")],
                     [("ЛС-0001", "", 1000)])
    r2 = c.post('/batch_calculate',
                data={'file': (BytesIO(raw2), 'r.xlsx'), 'calc_date': '2026-10-04'},
                content_type='multipart/form-data')
    print("  статус:", r2.status_code)
    assert r2.status_code == 400 and "дата платежа" in r2.data.decode()
    print("  ответ:", r2.data.decode()[:150] + "...")

    print()
    print("=== 3. Двойной вычет: по умолчанию предупреждение, в strict — стоп ===")
    raw3 = make_book([("ЛС-0001", "01.02.2026", 600, "15.03.2026", "ЖКУ")],
                     [("ЛС-0001", "15.04.2026", 2000)])
    r3 = c.post('/batch_calculate',
                data={'file': (BytesIO(raw3), 'r.xlsx'), 'calc_date': '2026-10-04',
                      'rate_mode': 'fixed95'},
                content_type='multipart/form-data')
    print("  без strict:", r3.status_code)
    assert r3.status_code == 200, "по умолчанию переплата не должна блокировать"
    wb3 = openpyxl.load_workbook(BytesIO(r3.data), data_only=True)
    for row in sheet_rows(wb3, "Сверка", header_row=4):
        if row.get("Лицевой счёт"):
            print(f"    статус сверки: {find(row, 'статус')}")

    r3s = c.post('/batch_calculate',
                 data={'file': (BytesIO(raw3), 'r.xlsx'), 'calc_date': '2026-10-04',
                       'strict': '1'},
                 content_type='multipart/form-data')
    print("  со strict:", r3s.status_code)
    assert r3s.status_code == 400
    print("    ответ:", r3s.data.decode()[:120] + "...")

print()
print("ВСЕ ПРОВЕРКИ ПРОЙДЕНЫ")
