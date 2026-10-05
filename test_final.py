#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Сверка трёх методов ставки на одном реестре, проверка аванса и контроля."""

import os
from io import BytesIO
import openpyxl

# Свежая база, чтобы прогон был воспроизводимым.
DB = os.path.join(os.path.dirname(os.path.abspath(__file__)), "peni.db")
for suffix in ("", "-journal", "-wal", "-shm"):
    try:
        os.remove(DB + suffix)
    except OSError:
        pass

from app import app, init_db                      # noqa: E402
from test_helpers import rows_of, find, total_row  # noqa: E402

init_db()


def book(charges, payments):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Начисления"
    ws.append(["Лицевой счёт", "Период", "Начисление", "Срок оплаты", "Вид долга"])
    for c in charges:
        ws.append(list(c))
    ws2 = wb.create_sheet("Платежи")
    ws2.append(["Лицевой счёт", "Дата платежа", "Сумма платежа"])
    for p in payments:
        ws2.append(list(p))
    b = BytesIO()
    wb.save(b)
    return b.getvalue()


MAIN = book([("ЛС-0001", "01.01.2026", 2500, "10.02.2026", "ЖКУ"),
             ("ЛС-0001", "01.02.2026", 2600, "15.03.2026", "ЖКУ"),
             ("ЛС-0002", "01.02.2026", 1800, "15.03.2026", "Капремонт"),
             ("ЛС-0002", "01.03.2026", 1800, "15.04.2026", "Капремонт"),
             ("ЛС-9568", "01.04.2021", 950, "10.05.2021", "ЖКУ")],
            [("ЛС-0001", "15.04.2026", 2500), ("ЛС-0002", "01.04.2026", 500),
             ("ЛС-9568", "01.01.2022", 900)])

ADVANCE = book([("ЛС-A", "01.02.2026", 4000, "15.03.2026", "ЖКУ")],
               [("ЛС-A", "01.03.2026", 1000)])      # взнос до срока оплаты

with app.test_client() as c:
    c.post('/login', data={'username': 'admin', 'password': 'admin123'})

    print("=== Три метода ставки на одном реестре ===")
    print(f"{'метод':18}{'пени':>12}{'всего':>14}  контроль")
    for mode in ("fixed95", "historical", "on_payment_date"):
        r = c.post('/batch_calculate',
                   data={'file': (BytesIO(MAIN), 'r.xlsx'), 'calc_date': '2026-10-04',
                         'rate_mode': mode},
                   content_type='multipart/form-data')
        assert r.status_code == 200, r.data.decode()[:200]
        wb = openpyxl.load_workbook(BytesIO(r.data), data_only=True)
        itog = total_row(wb, "Сводка")
        chk = total_row(wb, "Контроль", prefix="Контроль")
        print(f"{mode:18}{find(itog, 'пени'):>12,.2f}{find(itog, 'итого к оплате'):>14,.2f}"
              f"  {find(chk, 'статус')}")
        assert "сходятся" in str(find(chk, "статус")), "контроль не сошёлся"

    print()
    print("=== Аванс: взнос до срока оплаты ===")
    r = c.post('/batch_calculate',
               data={'file': (BytesIO(ADVANCE), 'a.xlsx'), 'calc_date': '2026-10-04',
                     'rate_mode': 'fixed95'},
               content_type='multipart/form-data')
    wb = openpyxl.load_workbook(BytesIO(r.data), data_only=True)
    row = rows_of(wb, "Сводка")[0]
    print(f"  начислено {find(row, 'начислено'):,.2f} | оплачено {find(row, 'оплачено'):,.2f}"
          f" | в т.ч. авансом {find(row, 'аванс'):,.2f}")
    print(f"  остаток {find(row, 'остаток'):,.2f} | пени {find(row, 'пени'):,.2f}")
    print(f"  примечание: {find(row, 'примечание')}")
    assert find(row, "аванс") == 1000.0
    assert find(row, "пени") < 500.0, "аванс должен снижать пени"

print()
print("ГОТОВО")
